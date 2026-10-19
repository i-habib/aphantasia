"""Render the P5 Brain Sciences manuscript and supplement as Word files (free-format MDPI submission).

Reads projects/P5_asd_adhd_replication/submission/brainsci/manuscript_brainsci.md and
submission/supplement.md, resolving [INCLUDE table.md] and [FIGURE file.png | caption] markers.
Requires python-docx. Run build_p5_submission.py first so tables and figures are current.
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SUBMISSION = ROOT / "projects/P5_asd_adhd_replication/submission"
OUT = SUBMISSION / "brainsci"
FONT = "Palatino Linotype"


def add_runs(paragraph, text, size=None):
    """Add text with **bold** and *italic* markup as separate runs."""
    for piece in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*)", text):
        if not piece:
            continue
        bold, italic = piece.startswith("**"), piece.startswith("*") and not piece.startswith("**")
        run = paragraph.add_run(piece.strip("*") if (bold or italic) else piece)
        run.bold, run.italic = bold or None, italic or None
        if size:
            run.font.size = Pt(size)


def add_table(document, block):
    rows = [line.strip().strip("|").split("|") for line in block.splitlines()
            if not set(line.replace("|", "").replace(":", "").replace(" ", "")) <= {"-"}]
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            add_runs(cell.paragraphs[0], value.strip(), size=8)
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    for row in table.rows:  # keep rows on one page where possible
        properties = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        properties.append(cant_split)
    document.add_paragraph()


def render(source_text, target, title_style=True):
    text = re.sub(r"^\[INCLUDE ([^\]]+)\]$", lambda m: (SUBMISSION / m.group(1)).read_text(), source_text, flags=re.M)
    document = Document()
    style = document.styles["Normal"]
    style.font.name, style.font.size = FONT, Pt(10)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        document.styles[name].font.name = FONT
    in_references = False
    for block in text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block.startswith("[FIGURE "):
            name, caption = [part.strip() for part in block[len("[FIGURE "):-1].split("|", 1)]
            document.add_picture(str(SUBMISSION / name), width=Inches(6.3))
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            label, _, rest = caption.partition(". ")
            paragraph = document.add_paragraph()
            add_runs(paragraph, f"**{label}.** {rest}", size=9)
            continue
        if block.startswith("|"):
            add_table(document, block)
            continue
        if block.startswith("# "):
            heading = document.add_heading(block[2:], level=0 if title_style else 1)
            continue
        if block.startswith("### "):
            document.add_heading(block[4:], level=2)
            continue
        if block.startswith("## "):
            document.add_heading(block[3:], level=1)
            in_references = block[3:].strip() == "References"
            continue
        if re.match(r"Table S?\d+\. ", block) and "\n" not in block and len(block) < 140:
            label, _, rest = block.partition(". ")
            paragraph = document.add_paragraph()
            add_runs(paragraph, f"**{label}.** {rest}", size=9)
            paragraph.paragraph_format.keep_with_next = True
            continue
        if re.match(r"(- |\d+\. )", block):
            for line in block.splitlines():
                marker = re.match(r"(\s*)(- |\d+\. )?(.*)", line)
                numbered = bool(marker.group(2)) and marker.group(2)[0].isdigit()
                if numbered and in_references:
                    paragraph = document.add_paragraph()
                    add_runs(paragraph, marker.group(2) + marker.group(3), size=9)
                    continue
                paragraph = document.add_paragraph(style="List Number" if numbered else "List Bullet")
                if marker.group(1):
                    paragraph.paragraph_format.left_indent = Inches(0.5)
                add_runs(paragraph, marker.group(3))
            continue
        paragraph = document.add_paragraph()
        add_runs(paragraph, block.replace("\n", " "))
    document.save(target)
    print(target)


def main():
    render((OUT / "manuscript_brainsci.md").read_text(), OUT / "manuscript_brainsci.docx")
    supplement = (SUBMISSION / "supplement.md").read_text().replace("# Supplementary material",
        "# Supplementary Materials: Autistic Adults More Often Report Absent Sensory Imagery", 1)
    render(supplement, OUT / "supplementary_materials.docx")
    render((OUT / "cover_letter_brainsci.md").read_text(), OUT / "cover_letter_brainsci.docx", title_style=False)


if __name__ == "__main__":
    main()
