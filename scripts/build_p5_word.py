"""Render the P5 Cortex manuscript and supplement as Word files.

Reads projects/P5_asd_adhd_replication/submission/cortex/manuscript_cortex.md and
submission/supplement.md, resolving [INCLUDE table.md] and [FIGURE file.png | caption] markers.
Requires python-docx. Run build_p5_submission.py first so tables and figures are current.
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SUBMISSION = ROOT / "projects/P5_asd_adhd_replication/submission"
OUT = SUBMISSION / "cortex"
FONT = "Times New Roman"


def add_runs(paragraph, text, size=None):
    """Add text with **bold** and *italic* markup as separate runs."""
    for piece in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*)", text):
        if not piece:
            continue
        bold, italic = piece.startswith("**"), piece.startswith("*") and not piece.startswith("**")
        run = paragraph.add_run(piece.strip("*") if (bold or italic) else piece)
        run.bold, run.italic = bold or None, italic or None
        if size:
            run.font.size = Pt(size)


def add_table(document, block):
    rows = [line.strip().strip("|").split("|") for line in block.splitlines()
            if not set(line.replace("|", "").replace(":", "").replace(" ", "")) <= {"-"}]
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            add_runs(cell.paragraphs[0], value.strip(), size=9)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.0
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    for row in table.rows:  # keep rows on one page where possible
        properties = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        properties.append(cant_split)


def render(source_text, target, title_style=True, line_numbers=True):
    text = re.sub(r"^\[INCLUDE ([^\]]+)\]$", lambda m: (SUBMISSION / m.group(1)).read_text(), source_text, flags=re.M)
    document = Document()
    # Standard manuscript layout: Times New Roman 12 pt, double-spaced, 1-inch margins,
    # black headings, continuous line numbers and page numbers.
    style = document.styles["Normal"]
    style.font.name, style.font.size = FONT, Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.paragraph_format.line_spacing = 2.0
    style.paragraph_format.space_after = Pt(0)
    for name, size in (("Title", 16), ("Heading 1", 13), ("Heading 2", 12), ("Heading 3", 12)):
        heading_style = document.styles[name]
        heading_style.font.name, heading_style.font.size = FONT, Pt(size)
        heading_style.font.bold = True
        heading_style.font.italic = name == "Heading 2"
        heading_style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = heading_style.element.rPr.rFonts
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attribute}"), None)
        heading_style.paragraph_format.space_before = Pt(12)
        heading_style.paragraph_format.space_after = Pt(0)
        heading_style.paragraph_format.line_spacing = 2.0
    document.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    border = document.styles["Title"].element.pPr.find(qn("w:pBdr"))  # default Title style has a rule under it
    if border is not None:
        document.styles["Title"].element.pPr.remove(border)
    section = document.sections[0]
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(1))
    if line_numbers:
        numbering = OxmlElement("w:lnNumType")
        numbering.set(qn("w:countBy"), "1")
        numbering.set(qn("w:restart"), "continuous")
        section._sectPr.append(numbering)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    title_page = "[PAGEBREAK]" in text
    in_references = False
    after_table = False
    for block in text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        note_follows_table, after_table = after_table, False
        if block == "[PAGEBREAK]":
            document.add_page_break()
            title_page = False
            continue
        if block.startswith("[FIGURE "):
            name, caption = [part.strip() for part in block[len("[FIGURE "):-1].split("|", 1)]
            document.add_picture(str(SUBMISSION / name), width=Inches(6.3))
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            label, _, rest = caption.partition(". ")
            paragraph = document.add_paragraph()
            add_runs(paragraph, f"**{label}.** {rest}", size=11)
            paragraph.paragraph_format.line_spacing = 1.15
            continue
        if block.startswith("|"):
            add_table(document, block)
            after_table = True
            continue
        if block.startswith("# "):
            heading = document.add_heading(block[2:], level=0 if title_style else 1)
            continue
        if block.startswith("### "):
            document.add_heading(block[4:], level=2)
            continue
        if block.startswith("## "):
            document.add_heading(block[3:], level=1)
            in_references = block[3:].strip() == "References"
            continue
        if re.match(r"Table S?\d+\. ", block) and "\n" not in block and len(block) < 140:
            label, _, rest = block.partition(". ")
            paragraph = document.add_paragraph()
            add_runs(paragraph, f"**{label}.** {rest}", size=9)
            paragraph.paragraph_format.keep_with_next = True
            continue
        if re.match(r"(- |\d+\. )", block):
            for line in block.splitlines():
                marker = re.match(r"(\s*)(- |\d+\. )?(.*)", line)
                numbered = bool(marker.group(2)) and marker.group(2)[0].isdigit()
                if numbered and in_references:
                    paragraph = document.add_paragraph()
                    add_runs(paragraph, marker.group(2) + marker.group(3), size=9)
                    continue
                paragraph = document.add_paragraph(style="List Number" if numbered else "List Bullet")
                if marker.group(1):
                    paragraph.paragraph_format.left_indent = Inches(0.5)
                add_runs(paragraph, marker.group(3))
            continue
        paragraph = document.add_paragraph()
        add_runs(paragraph, block.replace("\n", " "), size=10 if note_follows_table else None)
        if note_follows_table:  # table note: single-spaced, smaller, not indented
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(12)
        elif title_page:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif in_references:  # hanging indent, as in APA reference lists
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.5)
        elif not block.startswith("**"):
            paragraph.paragraph_format.first_line_indent = Inches(0.5)
    document.save(target)
    print(target)


def main():
    render((OUT / "manuscript_cortex.md").read_text(), OUT / "manuscript_cortex.docx")
    supplement = (SUBMISSION / "supplement.md").read_text().replace("# Supplementary material",
        "# Supplementary material: Autistic adults more often report absent sensory imagery", 1)
    render(supplement, OUT / "supplementary_material.docx", line_numbers=False)


if __name__ == "__main__":
    main()
